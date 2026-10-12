import io
import json
from contextlib import asynccontextmanager
from pathlib import Path

from apscheduler.schedulers.background import BackgroundScheduler
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from . import config, db, jobs

ESTADOS = {"nueva", "interesante", "descartada", "postulada"}
WEB = Path(__file__).resolve().parent / "web"

scheduler = BackgroundScheduler()


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init()
    scheduler.add_job(lambda: jobs.scrapear(), "interval", hours=config.SCRAPE_CADA_HORAS,
                      id="scrape", max_instances=1, coalesce=True)
    scheduler.start()
    yield
    scheduler.shutdown(wait=False)


app = FastAPI(title="Licitaciones SEACE - Fotocopiadoras", lifespan=lifespan)


def _where(q, estado, relevante, abiertas, entidad, vr_min, vr_max, desde):
    w, p = ["1=1"], []
    if q:
        w.append("(descripcion LIKE ? OR entidad LIKE ? OR nomenclatura LIKE ?)")
        p += [f"%{q}%"] * 3
    if estado:
        w.append("estado_seg = ?")
        p.append(estado)
    if relevante:
        w.append("relevante = 1")
    if abiertas:
        w.append("cierre_propuestas >= datetime('now','localtime')")
    if entidad:
        w.append("entidad LIKE ?")
        p.append(f"%{entidad}%")
    if vr_min is not None:
        w.append("valor_referencial >= ?")
        p.append(vr_min)
    if vr_max is not None:
        w.append("valor_referencial <= ?")
        p.append(vr_max)
    if desde:
        w.append("fecha_publicacion >= ?")
        p.append(desde)
    return " AND ".join(w), p


ORDEN = {
    "cierre": "(cierre_propuestas IS NULL), cierre_propuestas DESC",
    "cierre_asc": "(cierre_propuestas IS NULL), cierre_propuestas ASC",
    "publicacion": "fecha_publicacion DESC",
    "monto": "valor_referencial DESC",
    "entidad": "entidad ASC",
}


@app.get("/api/licitaciones")
def listar(q: str = "", estado: str = "", relevante: bool = True, abiertas: bool = False, entidad: str = "",
           vr_min: float | None = None, vr_max: float | None = None, desde: str = "",
           orden: str = "publicacion"):
    where, params = _where(q, estado, relevante, abiertas, entidad, vr_min, vr_max, desde)
    sql = f"""SELECT nid_proceso, nomenclatura, entidad, fecha_publicacion, objeto, descripcion,
                     valor_referencial, moneda, cierre_propuestas, estado_seg, relevante, first_seen,
                     ficha_ok, nota,
                     (cierre_propuestas >= datetime('now','localtime')) AS abierta,
                     (SELECT COUNT(*) FROM documentos d WHERE d.nid_proceso = l.nid_proceso) AS n_docs
              FROM licitaciones l WHERE {where} ORDER BY {ORDEN.get(orden, ORDEN['publicacion'])}"""
    with db.conn() as c:
        rows = [dict(r) for r in c.execute(sql, params)]
        stats = c.execute(
            """SELECT COUNT(*) total,
                      SUM(cierre_propuestas >= datetime('now','localtime')) abiertas,
                      SUM(estado_seg='nueva') nuevas,
                      SUM(estado_seg='interesante') interesantes
               FROM licitaciones WHERE relevante=1""").fetchone()
    return {"items": rows, "stats": dict(stats)}


@app.get("/api/licitaciones/{nid}")
def detalle(nid: str):
    with db.conn() as c:
        lic = c.execute("SELECT * FROM licitaciones WHERE nid_proceso=?", (nid,)).fetchone()
        if not lic:
            raise HTTPException(404, "No existe")
        crono = [dict(r) for r in c.execute(
            "SELECT etapa, inicio, fin FROM cronograma WHERE nid_proceso=? ORDER BY rowid", (nid,))]
        docs = [dict(r) for r in c.execute(
            "SELECT rowid AS id, etapa, documento, filename, tamano, fecha, local_path FROM documentos "
            "WHERE nid_proceso=? ORDER BY rowid", (nid,))]
    out = dict(lic)
    out["ficha"] = json.loads(out.pop("ficha_json") or "{}")
    out["cronograma"] = crono
    out["documentos"] = docs
    return out


class EstadoIn(BaseModel):
    estado: str | None = None
    nota: str | None = None


@app.post("/api/licitaciones/{nid}/estado")
def cambiar_estado(nid: str, body: EstadoIn):
    if body.estado is not None and body.estado not in ESTADOS:
        raise HTTPException(400, f"Estado invalido. Opciones: {sorted(ESTADOS)}")
    with db.conn() as c:
        if body.estado is not None:
            c.execute("UPDATE licitaciones SET estado_seg=? WHERE nid_proceso=?", (body.estado, nid))
        if body.nota is not None:
            c.execute("UPDATE licitaciones SET nota=? WHERE nid_proceso=?", (body.nota, nid))
    return {"ok": True}


@app.post("/api/scrape")
def lanzar_scrape(refrescar: bool = False):
    if not jobs.scrapear(refrescar=refrescar):
        raise HTTPException(409, "Ya hay una tarea en curso")
    return {"ok": True}


@app.post("/api/licitaciones/{nid}/descargar")
def descargar(nid: str):
    if not jobs.descargar(nid):
        raise HTTPException(409, "Ya hay una tarea en curso")
    return {"ok": True}


@app.get("/api/estado")
def estado_tareas():
    with db.conn() as c:
        corridas = [dict(r) for r in c.execute("SELECT * FROM corridas ORDER BY id DESC LIMIT 12")]
    job = scheduler.get_job("scrape")
    return {"tarea": jobs.estado, "log": list(jobs.logs)[-60:], "corridas": corridas,
            "proxima": str(job.next_run_time) if job and job.next_run_time else None}


@app.get("/api/documentos/{doc_id}")
def bajar_documento(doc_id: int):
    with db.conn() as c:
        r = c.execute("SELECT local_path, filename FROM documentos WHERE rowid=?", (doc_id,)).fetchone()
    if not r or not r["local_path"] or not Path(r["local_path"]).exists():
        raise HTTPException(404, "Documento no descargado todavia")
    return FileResponse(r["local_path"], filename=r["filename"])


@app.get("/api/export.xlsx")
def exportar(q: str = "", estado: str = "", relevante: bool = True, abiertas: bool = False, entidad: str = "",
             vr_min: float | None = None, vr_max: float | None = None, desde: str = ""):
    where, params = _where(q, estado, relevante, abiertas, entidad, vr_min, vr_max, desde)
    with db.conn() as c:
        rows = c.execute(
            f"SELECT * FROM licitaciones WHERE {where} ORDER BY fecha_publicacion DESC", params).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Licitaciones"
    cols = [("Nomenclatura", "nomenclatura", 30), ("Entidad", "entidad", 38), ("RUC", "ruc", 14),
            ("Publicada", "fecha_publicacion", 17), ("Cierre propuestas", "cierre_propuestas", 17),
            ("Objeto", "objeto", 11), ("Descripcion", "descripcion", 70), ("Valor referencial", "valor_referencial", 16),
            ("Moneda", "moneda", 9), ("Estado", "estado_seg", 13), ("Nota", "nota", 30)]
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C2410C")
    for r in rows:
        ws.append([r[c[1]] for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = c[2]
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
        row[7].number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="licitaciones_fotocopiadoras.xlsx"'})


app.mount("/", StaticFiles(directory=WEB, html=True), name="web")
